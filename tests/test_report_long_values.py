import base64
import io
from io import BytesIO

import openpyxl
from pypdf import PdfReader

from reports.artifacts import build_approved_pdf, build_evidence_xlsx
from reports.evidence import refresh_evidence_bundle
from reports.formatting import flatten_rows, shorten_for_report
from tests.test_project_client_and_export import _app, _call


def test_long_values_are_shortened_not_dumped_into_one_cell():
    b64 = base64.b64encode(b"x" * 60000).decode()
    assert "uploaded file content" in shorten_for_report("boq_xlsx_b64", b64)
    prose = "word " * 400
    short = shorten_for_report("notes", prose)
    assert len(short) < 700 and "characters in total" in short
    assert shorten_for_report("notes", "short text") == "short text"
    rows = flatten_rows({"boq_xlsx_b64": b64, "motor_kw": 5})
    assert all(len(str(v)) < 700 for _l, v, _u in rows)


def test_boq_job_with_uploaded_workbook_dispatches_a_pdf_and_xlsx():
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["Item", "Description", "Qty", "Unit"])
    for i in range(1, 400):
        ws.append([i, f"False ceiling gypsum board item number {i} with long tender description text", 10 + i, "Sqm"])
    buf = io.BytesIO(); wb.save(buf)
    b64 = base64.b64encode(buf.getvalue()).decode()
    assert len(b64) > 12000
    app = _app()
    st, job = _call(app, "POST", "/v1/jobs", {"source": "t", "requested_skill_id": "boq_takeoff", "inputs": {"boq_xlsx_b64": b64}})
    assert st.startswith("201"), job
    jid = job["job_id"]
    _call(app, "POST", f"/v1/jobs/{jid}/enqueue"); _call(app, "POST", f"/v1/jobs/{jid}/process"); _call(app, "POST", f"/v1/jobs/{jid}/approve", {"comment": "ok"})
    j = app.store.get(jid)
    bundle = refresh_evidence_bundle(job=j)
    pdf = build_approved_pdf(job=j, evidence_bundle=bundle)
    assert len(PdfReader(BytesIO(pdf)).pages) >= 1
    assert build_evidence_xlsx(job=j, evidence_bundle=bundle)


def test_boq_report_groups_materials_under_each_item():
    csv = ("item_no,description,unit,quantity\n"
           "1,Providing and fixing gypsum board partition 12.5mm both sides with GI studs,sqm,120\n"
           "2,Providing and fixing gypsum board false ceiling with GI grid,sqm,200\n"
           "3,Providing and fixing 12mm toughened glass partition with aluminium profile,sqm,40\n")
    app = _app()
    st, job = _call(app, "POST", "/v1/jobs", {"source": "t", "requested_skill_id": "boq_takeoff", "inputs": {"boq_csv": csv}})
    jid = job["job_id"]
    _call(app, "POST", f"/v1/jobs/{jid}/enqueue"); _call(app, "POST", f"/v1/jobs/{jid}/process"); _call(app, "POST", f"/v1/jobs/{jid}/approve", {"comment": "ok"})
    j = app.store.get(jid)
    bundle = refresh_evidence_bundle(job=j)
    pdf = build_approved_pdf(job=j, evidence_bundle=bundle)
    text = "\n".join(p.extract_text() for p in PdfReader(BytesIO(pdf)).pages)
    assert "Materials required for each BOQ item" in text and "Item 1" in text and "Item 2" in text
    wb = openpyxl.load_workbook(BytesIO(build_evidence_xlsx(job=j, evidence_bundle=bundle)))
    assert "BOQ Item Materials" in wb.sheetnames and "Line Materials" not in wb.sheetnames and "Line Takeoff" not in wb.sheetnames
    col = [r[0] for r in wb["BOQ Item Materials"].iter_rows(min_row=2, values_only=True) if r[0]]
    assert col[0].startswith("Item 1") and col[1].startswith("    ")
    assert "Indicative work programme" in text and "critical path" in text.lower()
    ws = wb["Work Programme"]
    assert ws.cell(1, 8).value == "Duration (days)" and ws.cell(2, 8).value >= 1
    open("/tmp/boq_sample.pdf", "wb").write(pdf)
