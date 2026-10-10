"""HVAC design package: room loads -> airflow -> duct sizes -> equipment schedule (preliminary)."""
import io
from types import SimpleNamespace

import pytest
from pypdf import PdfReader
import openpyxl

from reports.artifacts import build_approved_pdf, build_evidence_xlsx
from skills.common import SkillRequest
from skills.hvac.design_package import design_package
from skills.hvac.design_package.adapter import HVACDesignPackageSkill

ROOMS = [
    {"room_id": "R1", "name": "Open office", "area_m2": 120, "occupancy": 20},
    {"room_id": "R2", "name": "Meeting room", "area_m2": 30},
    {"room_id": "R3", "name": "Server room", "area_m2": 12},
]
INPUTS = {"building_type": "office", "climate_zone": "hot_dry", "rooms": ROOMS, "diversity_factor_pct": 90}


def test_package_is_consistent_and_labels_assumptions():
    r = design_package(building_type="office", climate_zone="hot_dry", rooms=ROOMS, diversity_factor_pct=90)
    sched = r["room_schedule"]
    assert [x["room_id"] for x in sched] == ["R1", "R2", "R3"]
    assert r["sum_of_room_loads_tr"] == pytest.approx(sum(x["cooling_load_tr"] for x in sched), abs=0.05)
    assert r["block_load_tr"] == pytest.approx(r["sum_of_room_loads_tr"] * 0.9, abs=0.05)
    for x in sched:
        assert x["supply_airflow_cfm"] == pytest.approx(x["cooling_load_tr"] * 400, abs=1)
        assert x["duct_diameter_mm"] >= 100 and 0 < x["duct_velocity_ms"] < 12
    names = {a["name"] for a in r["assumptions"]}
    assert {"cfm_per_tr", "target_velocity_ms", "diversity_factor_pct"} <= names
    assert any("Manual J" in s for s in r["limitations"])


@pytest.mark.parametrize("bad", [
    {"rooms": []},
    {"rooms": [{"room_id": "A"}]},
    {"rooms": [{"room_id": "A", "area_m2": -5}]},
    {"rooms": [{"room_id": "A", "area_m2": 10}, {"room_id": "A", "area_m2": 12}]},
    {"diversity_factor_pct": 0},
])
def test_invalid_inputs_are_rejected_not_guessed(bad):
    base = {"building_type": "office", "climate_zone": "hot_dry", "rooms": [{"room_id": "R1", "area_m2": 20}]}
    base.update(bad)
    with pytest.raises(ValueError):
        design_package(**base)


def test_adapter_reports_validation_errors_and_no_result():
    res = HVACDesignPackageSkill().run(SkillRequest(skill_id="hvac_design_package", inputs={"building_type": "office"}))
    assert res.status == "input_validation_failed" and not res.engineering_result
    assert any("climate_zone" in e for e in res.validation_errors) and any("rooms" in e for e in res.validation_errors)


def test_report_shows_schedule_as_table_in_pdf_and_excel():
    res = HVACDesignPackageSkill().run(SkillRequest(skill_id="hvac_design_package", inputs=INPUTS)).to_dict()
    assert res["status"] == "draft_ready" and res["human_review_required"] is True
    job = SimpleNamespace(job_id="j1", report_id="r1", status=SimpleNamespace(value="approved"),
                          requested_skill_id="hvac_design_package", skill_id="hvac_design_package", project_context={})
    bundle = {"manifest": {"schema_version": "engineering-evidence-v1", "bundle_sha256": "a", "input_hash": "b", "source_evidence": []},
              "request": {"source": "t", "inputs": INPUTS},
              "tasks": [{"task_id": "task-1", "capability_id": "hvac_design_package", "status": "completed", "objective": "HVAC design package",
                         "inputs": INPUTS, "engineering_result": res["engineering_result"], "calculation_trace": res["calculation_trace"],
                         "assumptions": res["assumptions"], "warnings": res["warnings"], "compliance": [], "skill_version": "1.0.0"}],
              "workflow": {"status": "completed", "blockers": []}, "result": {"status": "completed"}, "review": {"events": []}}
    text = " ".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(build_approved_pdf(job=job, evidence_bundle=bundle))).pages)
    assert "Room Schedule" in text and "Open office" in text and "Meeting room" in text and "Server room" in text
    assert "Chilled-water header size" in text
    wb = openpyxl.load_workbook(io.BytesIO(build_evidence_xlsx(job=job, evidence_bundle=bundle)))
    assert "Room Schedule" in wb.sheetnames
    rows = list(wb["Room Schedule"].iter_rows(values_only=True))
    assert len(rows) == 4 and any("Duct" in str(h) for h in rows[0])


def test_api_intake_review_approve_dispatch_for_design_package():
    import json
    from api.app import APIApp
    from crm.store import InMemoryCRMStore
    from integrations.service import InMemoryIntegrationStore
    from jobs.store import InMemoryJobStore

    app = APIApp(store=InMemoryJobStore(), crm_store=InMemoryCRMStore(), integration_store=InMemoryIntegrationStore())

    def call(method, path, body=None):
        raw = json.dumps(body or {}).encode()
        env = {"REQUEST_METHOD": method, "PATH_INFO": path, "CONTENT_LENGTH": str(len(raw)),
               "wsgi.input": io.BytesIO(raw), "HTTP_X_TENANT_ID": "tenant-a"}
        out = {}
        data = b"".join(app(env, lambda s, h: out.update(status=s)))
        return out["status"], json.loads(data.decode() or "{}")

    s, j = call("POST", "/v1/intake", {"message": "Prepare the HVAC design package for our office floor",
                                       "requested_skill_id": "hvac_design_package", "inputs": INPUTS})
    assert s.startswith("201") and j["job"]["status"] == "human_review", j
    s, j = call("POST", f"/v1/jobs/{j['job']['job_id']}/approve", {"comment": "ok", "dispatch": True})
    assert s.startswith("200") and j["status"] in ("dispatched", "completed")


def test_plant_summary_ventilation_and_fans_hand_checked():
    rooms = [{"room_id": "A", "name": "Open", "area_m2": 120, "occupancy": 20}, {"room_id": "B", "name": "Meet", "area_m2": 30}]
    r = design_package(building_type="office", climate_zone="hot_dry", rooms=rooms, outdoor_db_c=43, outdoor_rh_pct=30)
    a, b = r["room_schedule"]
    assert a["outdoor_air_l_s"] == 86.0            # 2.5 x 20 + 0.3 x 120
    assert b["outdoor_air_l_s"] == 9.0             # area component only, no occupancy given
    assert "only the area component" in b["outdoor_air_basis"]
    assert 3.2 < a["ventilation_load_kw"] < 3.9    # 0.086 m3/s x ~1.07 kg/m3 x ~38 kJ/kg
    ps = r["plant_summary"]
    kw = r["block_load_kw"]
    assert ps["chilled_water_flow_m3h"] == pytest.approx(kw / (4.18 * 5.5) * 3.6, rel=0.03)
    assert ps["chilled_water_header_dn"] >= 20
    assert ps["total_ventilation_load_kw_not_included"] > 0
    assert r["block_load_tr"] == pytest.approx(r["sum_of_room_loads_tr"], abs=0.01)  # ventilation not added


def test_ventilation_load_needs_both_outdoor_conditions():
    with pytest.raises(ValueError):
        design_package(building_type="office", climate_zone="hot_dry", rooms=[{"area_m2": 50}], outdoor_db_c=40)
