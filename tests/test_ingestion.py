import base64, io, json, zipfile
from pathlib import Path
from api.app import APIApp
from ingestion.extractors import extract_bytes, detect_source_type
from ingestion.service import IngestionService, LocalAttachmentStore


def make_xlsx_bytes():
    content_types = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/><Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/></Types>'''
    workbook = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?><workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets><sheet name="Inputs" sheetId="1" r:id="rId1"/></sheets></workbook>'''
    rels = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Target="worksheets/sheet1.xml" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet"/></Relationships>'''
    sheet = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?><worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>Airflow</t></is></c><c r="B1" t="inlineStr"><is><t>3600</t></is></c></row><row r="2"><c r="A2" t="inlineStr"><is><t>Velocity</t></is></c><c r="B2" t="inlineStr"><is><t>8</t></is></c></row></sheetData></worksheet>'''
    out = io.BytesIO()
    with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as z:
        z.writestr('[Content_Types].xml', content_types)
        z.writestr('xl/workbook.xml', workbook)
        z.writestr('xl/_rels/workbook.xml.rels', rels)
        z.writestr('xl/worksheets/sheet1.xml', sheet)
    return out.getvalue()


def test_source_type_detection():
    assert detect_source_type('spec.pdf') == 'pdf'
    assert detect_source_type('boq.xlsx') == 'xlsx'
    assert detect_source_type('plan.dxf') == 'dxf'
    assert detect_source_type('drawing.dwg') == 'cad_binary'


def test_text_and_csv_extraction():
    txt = extract_bytes(b'Project: Alpha\n\nAirflow: 3600 CFM', filename='note.txt', mime_type='text/plain', attachment_id='a1')
    assert txt.status == 'extracted' and 'Airflow' in txt.text and txt.chunks
    csv = extract_bytes(b'Item,Qty\nDuct,10\nDiffuser,4', filename='boq.csv', mime_type='text/csv', attachment_id='a2')
    assert csv.metadata['rows'] == 3 and 'Diffuser' in csv.text


def test_xlsx_extraction_without_openpyxl():
    result = extract_bytes(make_xlsx_bytes(), filename='inputs.xlsx', attachment_id='a3')
    assert result.status == 'extracted'
    assert result.metadata['worksheet_count'] == 1
    assert 'Airflow' in result.text and '3600' in result.text



def test_pdf_text_extraction():
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.drawString(72, 720, "Project Alpha - AHU Supply Air 14 C")
    c.save()
    result = extract_bytes(buf.getvalue(), filename='spec.pdf', mime_type='application/pdf', attachment_id='a5')
    assert result.status == 'extracted'
    assert result.metadata['pages'] == 1
    assert 'AHU Supply Air' in result.text


def test_docx_text_and_table_extraction():
    from docx import Document
    buf = io.BytesIO()
    doc = Document()
    doc.add_paragraph("Project Beta - Chilled Water System")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'Flow'
    table.cell(0, 1).text = '20 m3/hr'
    table.cell(1, 0).text = 'Head'
    table.cell(1, 1).text = '28 m'
    doc.save(buf)
    result = extract_bytes(buf.getvalue(), filename='scope.docx', mime_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document', attachment_id='a6')
    assert result.status == 'extracted'
    assert 'Project Beta' in result.text and '20 m3/hr' in result.text

def test_dxf_metadata_only():
    dxf = b'0\nSECTION\n2\nTABLES\n0\nTABLE\n2\nLAYER\n0\nLAYER\n2\nA-WALL\n0\nENDTAB\n0\nENDSEC\n0\nEOF\n'
    result = extract_bytes(dxf, filename='plan.dxf', attachment_id='a4')
    assert result.status == 'metadata_only'
    assert result.metadata['layers'] == ['A-WALL']


def test_ingestion_register_and_extract():
    service = IngestionService(LocalAttachmentStore())
    attachment = service.register(tenant_id='t1', job_id='j1', filename='note.txt', data=b'AHU-1 SAT = 14 C')
    result = service.extract(attachment)
    assert attachment.sha256 and result.status == 'extracted' and 'AHU-1' in result.text


def call(app, method, path, body=None, tenant='tenant-a'):
    raw=json.dumps(body or {}).encode()
    env={'REQUEST_METHOD':method,'PATH_INFO':path,'CONTENT_LENGTH':str(len(raw)),'wsgi.input':io.BytesIO(raw),'HTTP_X_TENANT_ID':tenant}
    result={}
    out=b''.join(app(env,lambda s,h: result.setdefault('status',s)))
    return result['status'], json.loads(out.decode())


def test_api_attachment_ingestion_flow():
    app=APIApp()
    status, job = call(app, 'POST', '/v1/jobs', {'source':'api','requested_skill_id':'duct_sizing','inputs':{}})
    assert status.startswith('201')
    jid=job['job_id']
    payload={'filename':'spec.txt','mime_type':'text/plain','content_base64':base64.b64encode(b'Airflow 3600 CFM').decode()}
    status, job = call(app, 'POST', f'/v1/jobs/{jid}/attachments', payload)
    assert status.startswith('201') and len(job['attachments']) == 1
    aid=job['attachments'][0]['attachment_id']
    status, result = call(app, 'POST', f'/v1/jobs/{jid}/attachments/{aid}', {})
    assert status.startswith('200') and result['status'] == 'extracted' and 'Airflow' in result['text']
