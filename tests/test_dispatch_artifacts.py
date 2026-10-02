from types import SimpleNamespace
from zipfile import ZipFile
from io import BytesIO

from reports.artifacts import build_approved_pdf, build_evidence_xlsx


def _job():
    return SimpleNamespace(
        job_id="job-artifact-1",
        report_id="report-1",
        status=SimpleNamespace(value="approved"),
        requested_skill_id="hvac_decarbonisation",
        skill_id="hvac_decarbonisation",
    )


def _bundle():
    return {
        "manifest": {
            "schema_version": "engineering-evidence-v1",
            "bundle_sha256": "abc123",
            "input_hash": "input",
            "source_evidence": [],
        },
        "request": {"source": "dashboard", "inputs": {"annual_energy_kwh": 100000}},
        "tasks": [{
            "task_id": "task-1",
            "capability_id": "hvac_decarbonisation",
            "status": "completed",
            "objective": "Assess energy savings",
            "inputs": {"annual_energy_kwh": 100000},
            "engineering_result": {"annual_saving_kwh": 8000},
            "calculation_trace": ["100000 * 0.08"],
            "assumptions": ["constant load"],
            "warnings": [],
            "compliance": [],
        }],
        "workflow": {"status": "completed", "blockers": []},
        "result": {"status": "completed"},
        "qa": {"status": "ready_for_human_review"},
        "governance": {},
        "review": {"required": True, "status": "approved", "events": []},
    }


def test_approved_pdf_is_watermarked_and_nonempty():
    data = build_approved_pdf(job=_job(), evidence_bundle=_bundle())
    assert data.startswith(b"%PDF")
    from pypdf import PdfReader
    text = "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(data)).pages)
    assert "APPROVED CONTROLLED DOCUMENT" in text


def test_evidence_workbook_is_valid_xlsx():
    data = build_evidence_xlsx(job=_job(), evidence_bundle=_bundle())
    assert data[:2] == b"PK"
    with ZipFile(BytesIO(data)) as archive:
        assert "xl/workbook.xml" in archive.namelist()
        workbook = archive.read("xl/workbook.xml")
        assert b"Tasks" in workbook
        assert b"Manifest" in workbook


def test_pdf_shows_tables_not_raw_json():
    from pypdf import PdfReader
    bundle = _bundle()
    bundle["tasks"][0]["engineering_result"] = {"annual_savings_kwh": 130233, "savings_pct": 28.9}
    data = build_approved_pdf(job=_job(), evidence_bundle=bundle)
    text = "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(data)).pages)
    assert "Annual Savings" in text and "130,233" in text and "kWh" in text
    assert "{" not in text and '":' not in text  # no raw JSON dumped into the report


def test_workbook_has_readable_result_sheets():
    from openpyxl import load_workbook
    wb = load_workbook(BytesIO(build_evidence_xlsx(job=_job(), evidence_bundle=_bundle())))
    for name in ("Summary", "Inputs", "Results", "Calculation Steps", "Assumptions & Notes", "Compliance"):
        assert name in wb.sheetnames, name
    rows = list(wb["Results"].iter_rows(values_only=True))
    assert rows[0] == ("Task", "Result", "Value", "Unit")
    assert any(r[1] == "Annual Saving" and r[2] == 8000 and r[3] == "kWh" for r in rows[1:])
