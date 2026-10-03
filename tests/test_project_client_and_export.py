import io
import json

from api.app import APIApp
from crm.store import InMemoryCRMStore
from integrations.service import InMemoryIntegrationStore
from jobs.store import InMemoryJobStore
from reports.artifacts import build_approved_pdf, build_evidence_xlsx
from reports.evidence import refresh_evidence_bundle
from reports.register import build_register_workbook

VFD = {"motor_kw": 75, "annual_hours": 6000, "tariff_per_kwh": 8.5, "speed_reduction_pct": 15, "static_head_fraction": 25}


def _app():
    return APIApp(store=InMemoryJobStore(), crm_store=InMemoryCRMStore(), integration_store=InMemoryIntegrationStore())


def _call(app, method, path, body=None):
    raw = json.dumps(body or {}).encode()
    env = {"REQUEST_METHOD": method, "PATH_INFO": path, "CONTENT_LENGTH": str(len(raw)),
           "wsgi.input": io.BytesIO(raw), "HTTP_X_TENANT_ID": "tenant-a"}
    out = {}
    data = b"".join(app(env, lambda s, h: out.update(status=s)))
    return out["status"], json.loads(data.decode() or "{}")


def _approved_job(app, *, project=None, client=None):
    payload = {"source": "t", "requested_skill_id": "vfd_energy_savings", "inputs": dict(VFD)}
    if project:
        payload["project"] = project
    if client:
        payload["client"] = client
    status, job = _call(app, "POST", "/v1/jobs", payload)
    assert status.startswith("201"), job
    jid = job["job_id"]
    _call(app, "POST", f"/v1/jobs/{jid}/enqueue")
    _call(app, "POST", f"/v1/jobs/{jid}/process")
    _call(app, "POST", f"/v1/jobs/{jid}/approve", {"comment": "ok"})
    return jid, app.store.get(jid)


def test_project_and_client_stored_on_job():
    app = _app()
    _, job = _approved_job(app, project="Hotel ABC Retrofit", client="Hotel ABC Pvt Ltd")
    assert job.project_context.get("project") == "Hotel ABC Retrofit"
    assert job.project_context.get("client") == "Hotel ABC Pvt Ltd"


def test_project_and_client_absent_when_not_supplied():
    app = _app()
    _, job = _approved_job(app)
    assert "project" not in job.project_context
    assert "client" not in job.project_context


def test_pdf_and_xlsx_show_project_and_client():
    from pypdf import PdfReader
    from io import BytesIO
    from openpyxl import load_workbook
    app = _app()
    _, job = _approved_job(app, project="Hotel ABC Retrofit", client="Hotel ABC Pvt Ltd")
    bundle = refresh_evidence_bundle(job=job)
    pdf_text = "\n".join(p.extract_text() or "" for p in PdfReader(BytesIO(build_approved_pdf(job=job, evidence_bundle=bundle))).pages)
    assert "Hotel ABC Retrofit" in pdf_text and "Hotel ABC Pvt Ltd" in pdf_text
    wb = load_workbook(BytesIO(build_evidence_xlsx(job=job, evidence_bundle=bundle)))
    summary_text = "\n".join(str(c) for row in wb["Summary"].iter_rows(values_only=True) for c in row)
    assert "Hotel ABC Retrofit" in summary_text and "Hotel ABC Pvt Ltd" in summary_text


def test_register_includes_project_and_client():
    from reports.register import rows_from_jobs
    app = _app()
    _, job = _approved_job(app, project="Hotel ABC Retrofit", client="Hotel ABC Pvt Ltd")
    rows = rows_from_jobs(app.store.list())
    row = next(r for r in rows if r["job_id"] == job.job_id)
    assert row["project"] == "Hotel ABC Retrofit"
    assert row["client"] == "Hotel ABC Pvt Ltd"


def test_register_export_endpoint_returns_a_real_workbook():
    from openpyxl import load_workbook
    from io import BytesIO
    app = _app()
    _approved_job(app, project="Hotel ABC Retrofit", client="Hotel ABC Pvt Ltd")
    env = {"REQUEST_METHOD": "GET", "PATH_INFO": "/v1/reports/export", "HTTP_X_TENANT_ID": "tenant-a"}
    out = {}
    data = b"".join(app(env, lambda s, h: out.update(status=s, headers=dict(h))))
    assert out["status"].startswith("200")
    assert "spreadsheet" in out["headers"]["Content-Type"]
    wb = load_workbook(BytesIO(data))
    rows = list(wb["Report Register"].iter_rows(values_only=True))
    assert rows[0][:6] == ("Date", "Job ID", "Report ID", "Revision", "Project", "Client")
    assert any("Hotel ABC Retrofit" in (r or ()) for r in rows[1:])


def test_register_workbook_builder_handles_empty_rows():
    from openpyxl import load_workbook
    from io import BytesIO
    wb = load_workbook(BytesIO(build_register_workbook([])))
    assert wb["Report Register"]["A1"].value == "Date"
