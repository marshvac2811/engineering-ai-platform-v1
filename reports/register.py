"""Central Engineering Report Register.

A flat, human-friendly index over every engineering report: one row per report
version with its job, status, reviewer, approval/dispatch dates, hashes and which
artifacts exist. Supabase remains the system of record; the register is a view.
"""
from __future__ import annotations

from typing import Any, Dict, Iterable, List, Optional

# View name -> job statuses included. "all" has no filter.
VIEWS: Dict[str, Optional[frozenset]] = {
    "all": None,
    "pending_review": frozenset({"draft_ready", "human_review"}),
    "approved": frozenset({"approved", "dispatching"}),
    "dispatched": frozenset({"dispatched"}),
    "completed": frozenset({"completed"}),
    "failed": frozenset({"failed", "retry", "cancelled", "rework"}),
}


def _s(value: Any) -> str:
    return "" if value is None else str(value)


def register_row(artifact: Dict[str, Any], job: Dict[str, Any]) -> Dict[str, Any]:
    """Merge one report-artifact row with its job summary into a register row."""
    job_status = _s(job.get("status")) or _s(artifact.get("status"))
    return {
        "report_id": artifact.get("report_id"),
        "job_id": artifact.get("job_id") or job.get("job_id"),
        "revision": artifact.get("version") or 1,
        "created_at": artifact.get("created_at") or job.get("created_at"),
        "title": artifact.get("title"),
        "request": job.get("request") or None,
        "project": (job.get("project_context") or {}).get("project") or None,
        "client": (job.get("project_context") or {}).get("client") or None,
        "skill_id": artifact.get("skill_id") or job.get("skill_id") or job.get("requested_skill_id"),
        "source": job.get("source"),
        "status": job_status,
        "report_status": artifact.get("status"),
        "reviewed_by": artifact.get("reviewer") or None,
        "approved_at": artifact.get("approved_at"),
        "dispatched_at": artifact.get("dispatched_at"),
        "pdf_available": bool(artifact.get("pdf_storage_path")),
        "xlsx_available": bool(artifact.get("xlsx_storage_path")),
        "pdf_filename": artifact.get("pdf_filename"),
        "pdf_sha256": artifact.get("pdf_sha256"),
        "xlsx_sha256": artifact.get("xlsx_sha256"),
        "evidence_sha256": artifact.get("evidence_sha256"),
    }


def rows_from_artifacts(artifacts: Iterable[Dict[str, Any]], jobs_by_id: Dict[str, Dict[str, Any]]) -> List[Dict[str, Any]]:
    rows = [register_row(a, jobs_by_id.get(str(a.get("job_id")), {})) for a in artifacts]
    rows.sort(key=lambda r: _s(r["created_at"]), reverse=True)
    latest: Dict[str, int] = {}
    for r in rows:
        key = _s(r["job_id"])
        latest[key] = max(latest.get(key, 0), int(r["revision"] or 1))
    for r in rows:  # an older revision of the same job is superseded by the newest one
        r["superseded"] = int(r["revision"] or 1) < latest[_s(r["job_id"])]
    return rows


def rows_from_jobs(jobs: Iterable[Any]) -> List[Dict[str, Any]]:
    """Fallback for stores without a report-artifact table (in-memory/local)."""
    rows = []
    for job in jobs:
        if not getattr(job, "report_id", None):
            continue
        events = list(getattr(job, "events", []) or [])
        approval = next((e for e in reversed(events) if _s(getattr(e, "status", "")) == "approved"), None)
        artifacts = ((getattr(job, "dispatch_result", None) or {}).get("artifacts") or {})
        pdf, xlsx = artifacts.get("pdf") or {}, artifacts.get("evidence_workbook") or {}
        status = job.status.value if hasattr(job.status, "value") else _s(job.status)
        dispatched = next((e.created_at for e in reversed(events) if _s(getattr(e, "status", "")) == "dispatched"), None)
        artifact = {
            "report_id": job.report_id, "job_id": job.job_id, "version": artifacts.get("version") or 1,
            "status": "approved" if status in {"approved", "dispatching", "dispatched", "completed"} else "draft",
            "title": f"Engineering Report - {job.skill_id or job.requested_skill_id or 'analysis'}",
            "skill_id": job.skill_id or job.requested_skill_id, "created_at": job.created_at,
            "reviewer": (getattr(approval, "metadata", None) or {}).get("reviewer") if approval else None,
            "approved_at": getattr(approval, "created_at", None) if approval else None,
            "dispatched_at": dispatched,
            "pdf_storage_path": pdf.get("storage_path"), "xlsx_storage_path": xlsx.get("storage_path"),
            "pdf_filename": pdf.get("filename"), "pdf_sha256": pdf.get("sha256"),
            "xlsx_sha256": xlsx.get("sha256"), "evidence_sha256": artifacts.get("evidence_sha256"),
        }
        request = ((getattr(job, "orchestration", None) or {}).get("normalized_request"))
        rows.append(register_row(artifact, {"job_id": job.job_id, "status": status, "source": job.source, "request": request,
                                            "skill_id": job.skill_id, "requested_skill_id": job.requested_skill_id,
                                            "created_at": job.created_at, "project_context": getattr(job, "project_context", {})}))
    rows.sort(key=lambda r: _s(r["created_at"]), reverse=True)
    return rows


def view_counts(rows: List[Dict[str, Any]]) -> Dict[str, int]:
    return {name: sum(1 for r in rows if statuses is None or r["status"] in statuses) for name, statuses in VIEWS.items()}


def filter_rows(rows: List[Dict[str, Any]], *, view: str = "all", q: str = "", skill: str = "") -> List[Dict[str, Any]]:
    if view not in VIEWS:
        raise ValueError(f"Unknown register view: {view}")
    statuses = VIEWS[view]
    needle = (q or "").strip().lower()
    out = []
    for r in rows:
        if statuses is not None and r["status"] not in statuses:
            continue
        if skill and _s(r["skill_id"]) != skill:
            continue
        if needle and needle not in " ".join(_s(r.get(k)) for k in ("report_id", "job_id", "title", "request", "skill_id", "reviewed_by", "status")).lower():
            continue
        out.append(r)
    return out


def build_register_workbook(rows: List[Dict[str, Any]]) -> bytes:
    """One spreadsheet of every report row, for the client's own reconciliation tracker."""
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Report Register"
    header = ["Date", "Job ID", "Report ID", "Revision", "Project", "Client", "Request", "Skill",
              "Status", "Reviewed By", "Approval Date", "Dispatch Date", "PDF Available", "Evidence XLSX Available",
              "PDF SHA-256", "Evidence SHA-256", "Superseded"]
    ws.append(header)
    fill, font = PatternFill("solid", fgColor="1F3A5F"), Font(bold=True, color="FFFFFF")
    for c in ws[1]:
        c.fill, c.font = fill, font
    for r in rows:
        ws.append([
            _s(r.get("created_at")), _s(r.get("job_id")), _s(r.get("report_id")), r.get("revision") or 1,
            _s(r.get("project")), _s(r.get("client")), _s(r.get("request")), _s(r.get("skill_id")),
            _s(r.get("status")), _s(r.get("reviewed_by")), _s(r.get("approved_at")), _s(r.get("dispatched_at")),
            "Yes" if r.get("pdf_available") else "No", "Yes" if r.get("xlsx_available") else "No",
            _s(r.get("pdf_sha256")), _s(r.get("evidence_sha256")), "Yes" if r.get("superseded") else "No",
        ])
    widths = [20, 24, 24, 10, 22, 22, 46, 22, 16, 22, 20, 20, 12, 18, 40, 40, 12]
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
