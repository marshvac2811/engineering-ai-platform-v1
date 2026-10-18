"""Final engineering artifact generation: controlled PDF and evidence workbook.

The PDF is the only client-facing dispatch artifact. The workbook is an
internal evidence/support artifact. Both are generated from the same governed
engineering evidence bundle.
"""
from __future__ import annotations

import hashlib
import io
import json
from datetime import datetime, timezone
from typing import Any, Dict

from openpyxl import Workbook
import math

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether

from reports.formatting import (
    assumption_parts, describe_item, extract_tables, flatten_rows, format_value, humanize, split_inputs_and_results, split_unit, trace_steps,
)


DEFAULT_WATERMARK = "MARS ENGINEERING SOLUTIONS • APPROVED CONTROLLED DOCUMENT"


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True, default=str)
    return str(value)


def _pdf_watermark(canvas, doc, watermark: str) -> None:
    """Per-page branding: faint monogram + diagonal watermark, branded header and address footer."""
    from branding import brand, draw_logo, footer_line
    b = brand()
    width, height = A4

    # Large faint monogram behind the content.
    draw_logo(canvas, width / 2 - 50 * mm, height / 2 - 50 * mm, 100 * mm, mono=True)

    # Diagonal text watermark.
    canvas.saveState()
    canvas.setFillColor(colors.Color(0.45, 0.45, 0.45, alpha=0.16))
    canvas.translate(width / 2, height / 2)
    canvas.rotate(32)
    top, _, rest = str(watermark).partition(" \u2022 ")
    canvas.setFont("Helvetica-Bold", 30)
    canvas.drawCentredString(0, 6, top)
    if rest:
        canvas.setFont("Helvetica-Bold", 16)
        canvas.drawCentredString(0, -18, rest)
    canvas.restoreState()

    # Header: logo, name, tagline, rule.
    draw_logo(canvas, 15 * mm, height - 22 * mm, 11 * mm)
    canvas.saveState()
    canvas.setFillColor(colors.HexColor("#14233c"))
    canvas.setFont("Helvetica-Bold", 11)
    canvas.drawString(29 * mm, height - 16.5 * mm, b["name"].upper())
    canvas.setFillColor(colors.HexColor("#555555"))
    canvas.setFont("Helvetica", 7)
    canvas.drawString(29 * mm, height - 20.5 * mm, b["tagline"])
    canvas.setFillColor(colors.HexColor("#c8442b"))
    canvas.setFont("Helvetica-Bold", 7)
    canvas.drawRightString(width - 15 * mm, height - 16.5 * mm, "CONTROLLED ENGINEERING RECORD")
    canvas.setStrokeColor(colors.HexColor("#c8442b"))
    canvas.setLineWidth(0.8)
    canvas.line(15 * mm, height - 24 * mm, width - 15 * mm, height - 24 * mm)
    canvas.restoreState()

    # Footer: address/contact, validity note, page number.
    canvas.saveState()
    canvas.setStrokeColor(colors.HexColor("#cccccc"))
    canvas.setLineWidth(0.4)
    canvas.line(15 * mm, 14 * mm, width - 15 * mm, 14 * mm)
    canvas.setFillColor(colors.HexColor("#14233c"))
    canvas.setFont("Helvetica", 7)
    canvas.drawCentredString(width / 2, 10.5 * mm, footer_line())
    canvas.setFillColor(colors.HexColor("#555555"))
    canvas.setFont("Helvetica", 6.5)
    canvas.drawString(15 * mm, 6.5 * mm, "Controlled engineering record \u2022 Human approval required for validity")
    canvas.drawRightString(width - 15 * mm, 6.5 * mm, f"Page {doc.page}")
    canvas.restoreState()


_ACCENT = colors.HexColor("#1f3a5f")
_STRIPE = colors.HexColor("#f2f5f9")
_APPROVED = {"approved", "dispatching", "dispatched", "completed"}


def _fmt_dt(value: Any) -> str:
    if not value:
        return "-"
    try:
        dt = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
        return dt.astimezone(timezone.utc).strftime("%d %b %Y, %H:%M UTC")
    except ValueError:
        return str(value)


def _approval(bundle: Dict[str, Any]) -> Dict[str, str]:
    for event in reversed((bundle.get("review") or {}).get("events") or []):
        if str(event.get("status")) == "approved":
            meta = event.get("metadata") or {}
            return {"by": str(meta.get("reviewer") or "-"), "comment": str(meta.get("comment") or "-"),
                    "at": _fmt_dt(event.get("created_at"))}
    return {"by": "-", "comment": "-", "at": "-"}


def _task_inputs(task: Dict[str, Any], request_inputs: Dict[str, Any]) -> Dict[str, Any]:
    inputs = task.get("inputs") or {}
    return inputs if inputs else (request_inputs or {})


# Narrative keys that have their own report sections (Assumptions / Points to note); showing them
# again as result rows duplicates them.
_NARRATIVE_KEYS = {"assumptions", "limitations", "warnings"}
# BOQ takeoff: the per-item material blocks replace these flat tables in the PDF and Excel.
_BOQ_FLAT_KEYS = {"line_takeoff", "line_materials", "boq_breakdown", "work_programme", "material_totals"}


def _boq_breakdown(task: Dict[str, Any]) -> list:
    """Per-BOQ-item blocks: each item with its own material list (SiteTrack-style)."""
    value = (task.get("engineering_result") or {}).get("boq_breakdown")
    return [b for b in value if isinstance(b, dict)] if isinstance(value, list) else []


def _material_totals(task: Dict[str, Any]) -> list:
    """Consolidated purchase list: every material summed across all BOQ items, grouped by trade."""
    value = (task.get("engineering_result") or {}).get("material_totals")
    rows = [r for r in value if isinstance(r, dict)] if isinstance(value, list) else []
    return sorted(rows, key=lambda r: (str(r.get("trade") or "General"), str(r.get("type") or "") != "Main material", str(r.get("material_name") or "").lower()))


def _programme_rows(task: Dict[str, Any]) -> list:
    value = (task.get("engineering_result") or {}).get("work_programme")
    return [r for r in value if isinstance(r, dict)] if isinstance(value, list) else []


def _qty_text(qty: Any, unit: Any) -> str:
    """Site-friendly quantity: whole numbers for counts, two decimals otherwise."""
    if not isinstance(qty, (int, float)):
        return str(qty if qty not in (None, "") else "-")
    if str(unit or "").strip().lower() in ("nos", "no", "pcs", "set", "sets", "each"):
        import math
        return f"{math.ceil(qty - 1e-9):,}"
    return f"{qty:,.2f}"


def _item_heading(b: Dict[str, Any]) -> str:
    qty = b.get("boq_quantity")
    qty_txt = f"{qty:,.2f}".rstrip("0").rstrip(".") if isinstance(qty, (int, float)) else str(qty or "")
    return f"Item {b.get('item_no')}  |  {qty_txt} {b.get('boq_unit') or ''}".strip()


def _plain_results(task: Dict[str, Any]) -> Dict[str, Any]:
    return {k: v for k, v in (task.get("engineering_result") or {}).items() if k not in _NARRATIVE_KEYS and k not in _BOQ_FLAT_KEYS}


def _clip(value: Any, limit: int = 150) -> str:
    """Long tender descriptions are shortened in the PDF table; the Excel evidence workbook keeps the full text."""
    text = str(value if value not in (None, "") else "-")
    return text if len(text) <= limit else text[: limit - 3].rstrip() + "..."


def _results_rows(task: Dict[str, Any], request_inputs: Dict[str, Any]):
    inputs, results = split_inputs_and_results(_task_inputs(task, request_inputs), _plain_results(task))
    results, _tables = extract_tables(results)
    inputs, _in_tables = extract_tables(inputs)
    return flatten_rows(inputs), flatten_rows(results)


def _input_tables(task: Dict[str, Any], request_inputs: Dict[str, Any]):
    inputs, _results = split_inputs_and_results(_task_inputs(task, request_inputs), _plain_results(task))
    return extract_tables(inputs)[1]


def _result_tables(task: Dict[str, Any], request_inputs: Dict[str, Any]):
    """Schedules (lists of flat records) in the task result, shown as real tables."""
    _inputs, results = split_inputs_and_results(_task_inputs(task, request_inputs), _plain_results(task))
    return extract_tables(results)[1]


def _has_money(rows) -> bool:
    return any(w in label.lower() for label, _v, _u in rows for w in ("cost", "tariff", "saving", "invest", "payback"))


def _compliance_rows(tasks) -> list:
    rows = []
    for t in tasks:
        for c in t.get("compliance") or []:
            if isinstance(c, dict):
                rows.append(c)
    return rows


def build_approved_pdf(*, job, evidence_bundle: Dict[str, Any], watermark: str = DEFAULT_WATERMARK) -> bytes:
    # The artifact is generated during DISPATCHING, but the controlled document
    # represents the already-approved engineering result. Keep lifecycle state in
    # the audit record while using the client-facing controlled-document label here.
    document_status = "Approved for controlled dispatch" if str(job.status.value) in _APPROVED else humanize(job.status.value)
    if str(job.status.value) not in _APPROVED:
        raise ValueError("Approved PDF can only be generated after human approval")

    out = io.BytesIO()
    base = getSampleStyleSheet()
    body = ParagraphStyle("Body", parent=base["BodyText"], fontSize=9, leading=12)
    small = ParagraphStyle("Small", parent=base["BodyText"], fontSize=8, leading=10)
    cell = ParagraphStyle("Cell", parent=base["BodyText"], fontSize=8.5, leading=11)
    cell_b = ParagraphStyle("CellB", parent=cell, fontName="Helvetica-Bold")
    head_cell = ParagraphStyle("HeadCell", parent=cell, fontName="Helvetica-Bold", textColor=colors.white)
    num = ParagraphStyle("Num", parent=cell, alignment=2)
    h1 = ParagraphStyle("H1", parent=base["Title"], textColor=_ACCENT, fontSize=22, spaceAfter=2)
    sub = ParagraphStyle("Sub", parent=base["Heading3"], textColor=colors.HexColor("#2e7d32"), spaceAfter=6)
    section = ParagraphStyle("Section", parent=base["Heading2"], textColor=_ACCENT, spaceBefore=12, spaceAfter=5)
    task_h = ParagraphStyle("TaskH", parent=base["Heading3"], textColor=_ACCENT, spaceBefore=8, spaceAfter=3)
    bullet = ParagraphStyle("Bullet", parent=body, leftIndent=10, bulletIndent=0)

    doc = SimpleDocTemplate(
        out, pagesize=A4, rightMargin=15 * mm, leftMargin=15 * mm, topMargin=29 * mm, bottomMargin=20 * mm,
        title=f"Engineering Report - {job.job_id}", author="Mars Engineering Solutions",
    )
    manifest = evidence_bundle.get("manifest", {})
    request = evidence_bundle.get("request", {})
    tasks = evidence_bundle.get("tasks", []) or []
    governance = evidence_bundle.get("governance", {}) or {}
    orch = request.get("orchestration") or {}
    request_inputs = request.get("inputs") or {}
    approval = _approval(evidence_bundle)
    width = 180 * mm

    def P(text, style=cell):
        return Paragraph(str(text if text not in (None, "") else "-").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"), style)

    def table(header, rows, widths, numeric_cols=()):
        data = [[Paragraph(h, head_cell) for h in header]]
        for r in rows:
            data.append([P(v, num if i in numeric_cols else cell) for i, v in enumerate(r)])
        t = Table(data, colWidths=widths, repeatRows=1)
        style = [
            ("BACKGROUND", (0, 0), (-1, 0), _ACCENT),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#b8c2cf")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]
        for i in range(1, len(data)):
            if i % 2 == 0:
                style.append(("BACKGROUND", (0, i), (-1, i), _STRIPE))
        t.setStyle(TableStyle(style))
        return t

    def wide_table(header, rows, widths):
        small_cell = ParagraphStyle("small_cell", parent=cell, fontSize=6.5, leading=8)
        small_head = ParagraphStyle("small_head", parent=head_cell, fontSize=6.5, leading=8)
        data = [[Paragraph(h, small_head) for h in header]]
        for r in rows:
            data.append([Paragraph(_clip(v).replace("&", "&amp;").replace("<", "&lt;"), small_cell) for v in r])
        t = Table(data, colWidths=widths, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), _ACCENT), ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#b8c2cf")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ]))
        return t

    esc_ = lambda v: str(v if v not in (None, "") else "-").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    def kv_table(rows):
        t = Table([[P(k, cell_b), P(v)] for k, v in rows], colWidths=[48 * mm, width - 48 * mm])
        t.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#b8c2cf")),
            ("BACKGROUND", (0, 0), (0, -1), _STRIPE), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        return t

    story = [Paragraph("Engineering Report", h1), Paragraph("Approved for controlled dispatch", sub)]
    capabilities = ", ".join(humanize(t.get("capability_id")) for t in tasks if t.get("capability_id")) or humanize(job.requested_skill_id or job.skill_id)
    project_context = getattr(job, "project_context", None) or {}
    header_rows = []
    if project_context.get("project"):
        header_rows.append(("Project", project_context.get("project")))
    if project_context.get("client"):
        header_rows.append(("Client", project_context.get("client")))
    header_rows += [
        ("Serial No.", getattr(job, "serial_no", None) or "-"), ("Report ID", job.report_id),
        ("Scope of work", capabilities),
        ("Request", orch.get("normalized_request") or (tasks[0].get("objective") if tasks else None)),
        ("Approved by", f"{approval['by']} on {approval['at']}"),
        ("Report date", _fmt_dt(datetime.now(timezone.utc).isoformat())),
        ("Report revision", str((getattr(job, "result", {}) or {}).get("report_revision") or (getattr(job, "orchestration", {}) or {}).get("report_revision") or 1)),
        ("Skill version", ", ".join(str(t.get("skill_version") or "-") for t in tasks) or "-"),
        ("Source revision", ", ".join(str(t.get("source_revision") or "-") for t in tasks) or "-"),
    ]
    story.append(kv_table(header_rows))

    # 1. Inputs
    story.append(Paragraph("1. Inputs Used", section))
    shown_inputs = False
    for index, task in enumerate(tasks, 1):
        in_rows, _ = _results_rows(task, request_inputs)
        if not in_rows and not _input_tables(task, request_inputs):
            continue
        shown_inputs = True
        if len(tasks) > 1:
            story.append(Paragraph(f"Task {index}: {P(humanize(task.get('capability_id')), body).text}", task_h))
        story.append(table(["Parameter", "Value", "Unit"], [(l, format_value(v), u) for l, v, u in in_rows],
                           [90 * mm, 50 * mm, 40 * mm], numeric_cols=(1,)))
        for tbl in _input_tables(task, request_inputs):
            story.append(Spacer(1, 4))
            story.append(Paragraph(P(tbl["title"], body).text, task_h))
            story.append(wide_table(tbl["headers"], tbl["rows"], [width / max(1, len(tbl["headers"]))] * len(tbl["headers"])))
        story.append(Spacer(1, 4))
    if not shown_inputs:
        story.append(Paragraph("No numeric inputs were recorded for this report.", body))
    docs = manifest.get("source_evidence") or []
    if docs:
        story.append(Paragraph("Source documents supplied", task_h))
        story.append(table(["File", "Type", "Extraction"], [(d.get("filename"), d.get("mime_type"), humanize(d.get("extraction_status"))) for d in docs],
                           [90 * mm, 50 * mm, 40 * mm]))

    # 2. Results
    story.append(Paragraph("2. Results", section))
    money_note = False
    for index, task in enumerate(tasks, 1):
        _, res_rows = _results_rows(task, request_inputs)
        title = task.get("objective") or humanize(task.get("capability_id"))
        block = [Paragraph(f"Task {index}: {P(title, body).text}", task_h),
                 Paragraph(f"Method: {P(humanize(task.get('capability_id')), small).text} &nbsp;|&nbsp; Status: {P(document_status, small).text}", small),
                 Spacer(1, 3)]
        story.append(KeepTogether(block))
        if res_rows:
            story.append(table(["Result", "Value", "Unit"], [(l, format_value(v), u) for l, v, u in res_rows],
                               [90 * mm, 50 * mm, 40 * mm], numeric_cols=(1,)))
            money_note = money_note or _has_money(res_rows)
        elif not _result_tables(task, request_inputs):
            story.append(Paragraph("No numeric results were produced for this task.", body))
        breakdown = _boq_breakdown(task)
        if breakdown:
            story.append(Paragraph("Materials required for each BOQ item", task_h))
            story.append(Paragraph("Each BOQ description is followed by the materials it needs (quantity includes wastage). Combined purchase quantities are in the Material Totals table.", small))
            story.append(Spacer(1, 3))
            item_head = ParagraphStyle("item_head", parent=cell_b, fontSize=8, leading=10, textColor=colors.white)
            item_desc = ParagraphStyle("item_desc", parent=cell, fontSize=7.5, leading=9.5)
            sub_head = ParagraphStyle("sub_head", parent=cell_b, fontSize=7.5, leading=9)
            mat_cell = ParagraphStyle("mat_cell", parent=cell, fontSize=7.5, leading=9.5)
            mat_num = ParagraphStyle("mat_num", parent=mat_cell, alignment=2)
            esc = lambda v: str(v if v not in (None, "") else "-").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            for b in breakdown:
                note_ = b.get("drawing_note")
                data = [[Paragraph(esc(_item_heading(b)), item_head), "", "", ""],
                        [Paragraph(esc(_clip(b.get("description"), 400)), item_desc), "", "", ""]]
                if note_:
                    data.append([Paragraph("<b>From drawing:</b> " + esc(note_), item_desc), "", "", ""])
                data += [
                        [Paragraph(h, sub_head) for h in ("Material", "Unit", "Qty", "Basis")]]
                for m in b.get("materials") or []:
                    qty = m.get("quantity")
                    data.append([Paragraph(esc(m.get("material_name")), mat_cell), Paragraph(esc(m.get("unit")), mat_cell),
                                 Paragraph(esc(_qty_text(qty, m.get("unit"))), mat_num), Paragraph(esc(m.get("basis")), ParagraphStyle("b", parent=mat_cell, fontSize=6.5, leading=8))])
                t = Table(data, colWidths=[72 * mm, 14 * mm, 20 * mm, 74 * mm], repeatRows=0)
                hdr_row = 3 if note_ else 2
                t.setStyle(TableStyle([
                    ("SPAN", (0, 0), (-1, 0)), ("SPAN", (0, 1), (-1, 1))] + ([("SPAN", (0, 2), (-1, 2)), ("BACKGROUND", (0, 2), (-1, 2), colors.HexColor("#fff4d6"))] if note_ else []) + [
                    ("BACKGROUND", (0, 0), (-1, 0), _ACCENT), ("BACKGROUND", (0, 1), (-1, 1), _STRIPE),
                    ("BACKGROUND", (0, hdr_row), (-1, hdr_row), colors.HexColor("#dde4ee")),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#b8c2cf")),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
                ]))
                story.append(KeepTogether([t, Spacer(1, 6)]) if len(data) <= 22 else t)
        totals_ = _material_totals(task)
        if totals_:
            story.append(PageBreak())   # the consolidated list is a stand-alone page for purchasing
            story.append(Paragraph("Consolidated material list (all BOQ items together)", task_h))
            story.append(Paragraph("Each material is added up across every BOQ item, grouped by trade, so it can be used as one purchase list. Quantities include wastage. Supporting materials (fasteners, adhesives, tapes, sealants and similar) use standard starter factors to be confirmed by the engineer.", small))
            story.append(Spacer(1, 3))
            c_cell = ParagraphStyle("c_cell", parent=cell, fontSize=7.5, leading=9.5)
            c_num = ParagraphStyle("c_num", parent=c_cell, alignment=2)
            c_head = ParagraphStyle("c_head", parent=cell_b, fontSize=7.5, leading=9, textColor=colors.white)
            c_trade = ParagraphStyle("c_trade", parent=cell_b, fontSize=8, leading=10)
            esc3 = lambda v: str(v if v not in (None, "") else "-").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            data = [[Paragraph(h, c_head) for h in ("No.", "Material", "Unit", "Total qty", "Type", "BOQ items")]]
            style = [("BACKGROUND", (0, 0), (-1, 0), _ACCENT), ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#b8c2cf")),
                     ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]
            last_trade, n = None, 0
            for r in totals_:
                trade = str(r.get("trade") or "General")
                if trade != last_trade:
                    data.append([Paragraph(esc3(trade.upper()), c_trade), "", "", "", "", ""])
                    style += [("SPAN", (0, len(data) - 1), (-1, len(data) - 1)), ("BACKGROUND", (0, len(data) - 1), (-1, len(data) - 1), colors.HexColor("#dde4ee")),
                              ("NOSPLIT", (0, len(data) - 1), (-1, len(data)))]   # keep a trade heading with its first material row
                    last_trade = trade
                n += 1
                data.append([Paragraph(str(n), c_cell), Paragraph(esc3(_clip(r.get("material_name"), 90)), c_cell), Paragraph(esc3(r.get("unit")), c_cell),
                             Paragraph(esc3(_qty_text(r.get("quantity"), r.get("unit"))), c_num),
                             Paragraph("Main" if r.get("type") == "Main material" else "Supporting", c_cell),
                             Paragraph(esc3(_clip(r.get("source_items"), 60)), ParagraphStyle("c_src", parent=c_cell, fontSize=6.5, leading=8))])
            ct = Table(data, colWidths=[9 * mm, 66 * mm, 13 * mm, 22 * mm, 20 * mm, 50 * mm], repeatRows=1)
            ct.setStyle(TableStyle(style))
            story.append(ct)
            story.append(Spacer(1, 6))
        prog = _programme_rows(task)
        if prog:
            res_ = (task.get("engineering_result") or {})
            story.append(Paragraph("Indicative work programme", task_h))
            summ = f"Estimated site duration: {res_.get('total_duration_days')} working days (about {res_.get('total_duration_weeks')} weeks, {res_.get('programme_working_days_per_week')}-day week)"
            if res_.get("programme_finish_date"):
                summ += f", finishing {res_.get('programme_finish_date')}"
            story.append(Paragraph(P(summ + ".", body).text, body))
            story.append(Spacer(1, 3))
            total_days = int(res_.get("total_duration_days") or 0)
            wd_ = int(res_.get("programme_working_days_per_week") or 6)
            plen = 1 if total_days <= 36 else wd_          # one column per day for short jobs, per week otherwise
            unit_name = "day" if plen == 1 else "week"
            n_weeks = max(1, math.ceil(total_days / plen))
            bucket = max(1, math.ceil(n_weeks / 18))
            n_cols = math.ceil(n_weeks / bucket)
            story.append(Paragraph(f"Work is sequenced by stage; stages overlap where different gangs can work in different areas at the same time. Column numbers are {unit_name}s (each column is {bucket} {unit_name}{'s' if bucket > 1 else ''}). Red bars are the critical path: a delay there delays handover. This is an indicative programme for planning, to be confirmed with the site team.", small))
            g_small = ParagraphStyle("g_small", parent=cell, fontSize=6.5, leading=8)
            g_head = ParagraphStyle("g_head", parent=head_cell, fontSize=6.5, leading=8)
            head = [Paragraph(h, g_head) for h in ("No.", "Work", "Trade", "Days", "From", "To")]
            head += [Paragraph(str(c * bucket + 1), ParagraphStyle("g_wk", parent=g_head, fontSize=5.5, leading=6.5, alignment=1)) for c in range(n_cols)]
            data = [head]
            for r in prog:
                data.append([Paragraph(esc_(r.get("item_no")), g_small), Paragraph(esc_(_clip(r.get("description"), 70)), g_small), Paragraph(esc_(r.get("trade")), g_small),
                             Paragraph(esc_(r.get("duration_days")), g_small), Paragraph(esc_(r.get("start_day")), g_small), Paragraph(esc_(r.get("end_day")), g_small)] + [""] * n_cols)
            fixed = [12 * mm, 46 * mm, 19 * mm, 11 * mm, 11 * mm, 11 * mm]
            col_w = max(3.5 * mm, (width - sum(fixed)) / n_cols)
            gt = Table(data, colWidths=fixed + [col_w] * n_cols, repeatRows=1)
            style = [("BACKGROUND", (0, 0), (-1, 0), _ACCENT), ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#b8c2cf")),
                     ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
                     ("LEFTPADDING", (6, 0), (-1, -1), 0), ("RIGHTPADDING", (6, 0), (-1, -1), 0)]
            for i, r in enumerate(prog, 1):
                w0, w1 = (int(r["start_day"]) - 1) // plen, (int(r["end_day"]) - 1) // plen
                for c in range(n_cols):
                    if w0 <= (c + 1) * bucket - 1 and w1 >= c * bucket:
                        style.append(("BACKGROUND", (6 + c, i), (6 + c, i), colors.HexColor("#c0392b") if r.get("critical_path") == "Yes" else colors.HexColor("#4a78b5")))
            gt.setStyle(TableStyle(style))
            story.append(gt)
            story.append(Spacer(1, 6))
        for tbl in _result_tables(task, request_inputs):
            drop_prefixes = ("source calculation", "outdoor air basis", "load method", "cooling load (kw)", "duct friction", "room id")
            skip = {i for i, h in enumerate(tbl["headers"]) if h.lower().startswith(drop_prefixes)}
            basis_idx = next((i for i, h in enumerate(tbl["headers"]) if h.lower().startswith("outdoor air basis")), None)
            heads = [h for i, h in enumerate(tbl["headers"]) if i not in skip]
            rows_ = [[v for i, v in enumerate(r) if i not in skip] for r in tbl["rows"]]
            story.append(Paragraph(P(tbl["title"], body).text, task_h))
            # column widths follow content length so names and equipment text are not squeezed into a few letters per line
            weights = [max(9, min(26, max([len(str(h))] + [len(str(r[j])) for r in rows_]))) for j, h in enumerate(heads)]
            weights[0] = max(weights[0], 13)
            total_w = float(sum(weights))
            story.append(wide_table(heads, rows_, [width * w / total_w for w in weights]))
            if basis_idx is not None:
                bases = []
                for r in tbl["rows"]:
                    if r[basis_idx] not in bases and r[basis_idx] not in ("", "-"):
                        bases.append(r[basis_idx])
                for b_ in bases:
                    story.append(Paragraph("Outdoor air basis: " + str(b_).replace("&", "&amp;").replace("<", "&lt;"), small))
            story.append(Spacer(1, 4))
        steps = trace_steps(task.get("calculation_trace") or [])
        if steps:
            story.append(Paragraph("How the result was calculated", task_h))
            story.append(table(["Step", "Calculation step", "Detail"], [(x["step"], x["operation"], x["detail"]) for x in steps],
                               [15 * mm, 100 * mm, 65 * mm]))
        assumptions = task.get("assumptions") or []
        if assumptions:
            story.append(Paragraph("Assumptions", task_h))
            parts = [assumption_parts(a) for a in assumptions]
            if all(v for _l, v, _b in parts):
                story.append(table(["Assumption", "Value used", "Basis"], parts, [52 * mm, 38 * mm, 90 * mm]))
            else:
                story.extend(Paragraph(describe_item(a).replace("&", "&amp;").replace("<", "&lt;"), bullet, bulletText="•") for a in assumptions)
        warnings = task.get("warnings") or []
        if warnings:
            story.append(Paragraph("Points to note", task_h))
            story.extend(Paragraph(describe_item(w).replace("&", "&amp;").replace("<", "&lt;"), bullet, bulletText="•") for w in warnings)
        story.append(Spacer(1, 6))
    if money_note:
        story.append(Paragraph("Monetary values are in the currency of the tariff or cost figures supplied in the inputs.", small))

    # 3. Standards and compliance
    story.append(Paragraph("3. Standards and Compliance", section))
    limitations = []
    already_shown = {str(describe_item(w)) for t in tasks for w in (t.get("warnings") or [])}
    for t in tasks:
        for item in t.get("limitations") or []:
            text = str(item)
            if text and text not in limitations and text not in already_shown:   # "Points to note" already lists these
                limitations.append(text)
    if limitations:
        story.append(Paragraph("Limitations", task_h))
        story.extend(Paragraph(str(x).replace("&", "&amp;").replace("<", "&lt;"), bullet, bulletText="•") for x in limitations)


    checks = _compliance_rows(tasks)
    if checks:
        story.append(table(["Authority", "Code / edition", "Clause", "Requirement", "Project value", "Status"],
                           [(c.get("authority"), f"{c.get('code') or c.get('code_name') or ''} {c.get('edition') or ''}".strip(), c.get("clause"),
                             c.get("requirement"), format_value(c.get("project_value")), humanize(c.get("status"))) for c in checks],
                           [25 * mm, 32 * mm, 20 * mm, 43 * mm, 30 * mm, 30 * mm]))
    else:
        allowed = (governance.get("applicability") or {}).get("compliance_claim_allowed")
        story.append(Paragraph(
            "This report presents engineering calculation results only. No compliance claim is made, because verified "
            "standards evidence and project applicability have not been established." if not allowed else
            "Compliance claims are supported by the verified evidence recorded in the evidence workbook.", body))
        cands = governance.get("candidate_sources") or []
        if cands:
            story.append(Paragraph("Reference standards that may apply (not yet verified for this project)", task_h))
            story.append(table(["Standard", "Authority", "Edition"], [(c.get("title"), c.get("authority"), c.get("edition") or "-") for c in cands],
                               [100 * mm, 40 * mm, 40 * mm]))

    # 4. Review
    story.append(Paragraph("4. Review and Approval", section))
    story.append(kv_table([("Human review", "Required, completed"), ("Approved by", approval["by"]),
                           ("Approval date", approval["at"]), ("Reviewer comment", approval["comment"])]))

    # Appendix
    story.append(Paragraph("Appendix: Traceability", section))
    story.append(Paragraph("These identifiers allow the report to be matched to the stored engineering evidence record.", small))
    story.append(Spacer(1, 3))
    story.append(kv_table([
        ("Serial No.", getattr(job, "serial_no", None) or "-"), ("Job ID", job.job_id), ("Report ID", job.report_id),
        ("Evidence record SHA-256", manifest.get("bundle_sha256")), ("Inputs SHA-256", manifest.get("input_hash")),
        ("Evidence schema", manifest.get("schema_version")),
    ]))
    story.append(Spacer(1, 4))
    story.append(Paragraph("This document is generated from the platform evidence record, which is authoritative for traceability. "
                           "Human approval is required for validity.", small))

    doc.build(story,
              onFirstPage=lambda c, d: _pdf_watermark(c, d, watermark),
              onLaterPages=lambda c, d: _pdf_watermark(c, d, watermark))
    return out.getvalue()


def build_evidence_xlsx(*, job, evidence_bundle: Dict[str, Any]) -> bytes:
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    header_fill = PatternFill("solid", fgColor="1F3A5F")
    header_font = Font(bold=True, color="FFFFFF")
    wrap = Alignment(wrap_text=True, vertical="top")
    manifest = evidence_bundle.get("manifest", {})
    request = evidence_bundle.get("request", {})
    tasks = evidence_bundle.get("tasks", []) or []
    # Recover task outputs from the consolidated result for older persisted evidence records.
    if not tasks:
        consolidated = evidence_bundle.get("result") or {}
        engineering = consolidated.get("engineering_result") or {}
        recovered = engineering.get("task_results") or engineering.get("results") or []
        for item in recovered:
            if not isinstance(item, dict):
                continue
            tasks.append({
                "task_id": item.get("task_id"), "capability_id": item.get("capability_id"),
                "status": item.get("status"), "objective": item.get("objective"),
                "inputs": item.get("inputs") or request.get("inputs") or {},
                "engineering_result": item.get("engineering_result") or {},
                "calculation_trace": item.get("calculation_trace") or [],
                "assumptions": item.get("assumptions") or [], "warnings": item.get("warnings") or [],
                "compliance": item.get("compliance") or [], "skill_version": item.get("skill_version"),
                "source_revision": item.get("source_revision"), "limitations": item.get("limitations") or [],
            })
    governance = evidence_bundle.get("governance", {}) or {}
    request_inputs = request.get("inputs") or {}
    approval = _approval(evidence_bundle)

    def styled(sheet, header, rows, widths, number_cols=()):
        sheet.append(header)
        for c in sheet[1]:
            c.fill, c.font, c.alignment = header_fill, header_font, wrap
        for r in rows:
            sheet.append(list(r))
        for i, w in enumerate(widths):
            sheet.column_dimensions[chr(65 + i)].width = w
        for row in sheet.iter_rows(min_row=2):
            for c in row:
                c.alignment = wrap
                if c.column - 1 in number_cols and isinstance(c.value, (int, float)):
                    c.number_format = "#,##0.###"
        sheet.freeze_panes = "A2"

    # Summary (also keeps the legacy "Cover" fields)
    ws = wb.active
    ws.title = "Summary"
    document_status = "Approved for controlled dispatch" if str(job.status.value) in _APPROVED else humanize(job.status.value)
    project_context = getattr(job, "project_context", None) or {}
    summary_rows = []
    if project_context.get("project"):
        summary_rows.append(("Project", project_context.get("project")))
    if project_context.get("client"):
        summary_rows.append(("Client", project_context.get("client")))
    summary_rows += [
        ("Serial No.", getattr(job, "serial_no", None) or ""), ("Job ID", job.job_id), ("Report ID", job.report_id), ("Status", document_status),
        ("Scope of work", ", ".join(humanize(t.get("capability_id")) for t in tasks if t.get("capability_id")) or humanize(job.requested_skill_id or job.skill_id)),
        ("Request", (request.get("orchestration") or {}).get("normalized_request")),
        ("Approved by", approval["by"]), ("Approval date", approval["at"]),
        ("Evidence record SHA-256", manifest.get("bundle_sha256")), ("Evidence schema", manifest.get("schema_version")),
        ("Skill version", ", ".join(str(t.get("skill_version") or "-") for t in tasks) or "-"),
        ("Source revision", ", ".join(str(t.get("source_revision") or "-") for t in tasks) or "-"),
        ("Limitations", "; ".join(str(x) for t in tasks for x in ((t.get("limitations") or []) + (t.get("warnings") or []))) or "-"),
        ("Generated", _fmt_dt(datetime.now(timezone.utc).isoformat())),
    ]
    styled(ws, ["Item", "Detail"], summary_rows, [30, 100])

    inputs_rows, results_rows, step_rows, note_rows, check_rows = [], [], [], [], []
    for index, t in enumerate(tasks, 1):
        label = f"Task {index}: {humanize(t.get('capability_id'))}"
        in_rows, res_rows = _results_rows(t, request_inputs)
        inputs_rows += [(label, l, v, u) for l, v, u in in_rows]
        results_rows += [(label, l, v, u) for l, v, u in res_rows]
        step_rows += [(label, x["step"], x["operation"], x["detail"]) for x in trace_steps(t.get("calculation_trace") or [])]
        note_rows += [(label, "Assumption", (lambda p: f"{p[0]} = {p[1]}. {p[2]}".strip() if p[1] else p[0])(assumption_parts(a))) for a in t.get("assumptions") or []]
        note_rows += [(label, "Point to note", describe_item(w)) for w in t.get("warnings") or []]
        for c in t.get("compliance") or []:
            if isinstance(c, dict):
                check_rows.append((label, c.get("authority"), f"{c.get('code') or c.get('code_name') or ''} {c.get('edition') or ''}".strip(),
                                   c.get("clause"), c.get("requirement"), c.get("project_value"), c.get("required_value"), humanize(c.get("status"))))

    styled(wb.create_sheet("Inputs"), ["Task", "Parameter", "Value", "Unit"], inputs_rows, [30, 40, 18, 16], number_cols=(2,))
    styled(wb.create_sheet("Results"), ["Task", "Result", "Value", "Unit"], results_rows, [30, 40, 18, 16], number_cols=(2,))
    for index, t in enumerate(tasks, 1):
        for tbl in _input_tables(t, request_inputs) + _result_tables(t, request_inputs):
            sheet_name = f"{tbl['title']}"[:28] or "Schedule"
            styled(wb.create_sheet(sheet_name), tbl["headers"], tbl["rows"], [max(14, min(30, len(h) + 4)) for h in tbl["headers"]])
    for index, t in enumerate(tasks, 1):
        breakdown = _boq_breakdown(t)
        if not breakdown:
            continue
        sh = wb.create_sheet("BOQ Item Materials")
        sh.append(["Item / Material", "Unit", "Quantity (incl. wastage)", "Basis"])
        for c in sh[1]:
            c.fill, c.font, c.alignment = header_fill, header_font, wrap
        item_fill = PatternFill("solid", fgColor="D9E1F2")
        for b in breakdown:
            sh.append([f"{_item_heading(b)}  -  {b.get('description') or ''}", b.get("boq_unit"), b.get("boq_quantity"), b.get("recipe_name")])
            r = sh.max_row
            for c in sh[r]:
                c.fill, c.font, c.alignment = item_fill, Font(bold=True), wrap
            if b.get("drawing_note"):
                sh.append(["    From drawing: " + str(b["drawing_note"])])
                sh.cell(sh.max_row, 1).font = Font(italic=True, color="7A5200")
                sh.cell(sh.max_row, 1).alignment = wrap
            for m in b.get("materials") or []:
                sh.append(["    " + str(m.get("material_name")), m.get("unit"), (round(m["quantity"], 2) if isinstance(m.get("quantity"), (int, float)) else m.get("quantity")), m.get("basis")])
                for c in sh[sh.max_row]:
                    c.alignment = wrap
                sh.cell(sh.max_row, 3).number_format = "#,##0.00"
            sh.append([])
        for col, w in zip("ABCD", (80, 12, 22, 60)):
            sh.column_dimensions[col].width = w
        sh.freeze_panes = "A2"
        break
    for index, t in enumerate(tasks, 1):
        totals_ = _material_totals(t)
        if not totals_:
            continue
        sh = wb.create_sheet("Consolidated Materials")
        sh.append(["No.", "Trade", "Material", "Unit", "Total quantity (incl. wastage)", "Type", "BOQ items using it"])
        for c in sh[1]:
            c.fill, c.font, c.alignment = header_fill, header_font, wrap
        for n, r in enumerate(totals_, 1):
            q = r.get("quantity")
            sh.append([n, r.get("trade") or "General", r.get("material_name"), r.get("unit"), (round(q, 2) if isinstance(q, (int, float)) else q),
                       "Main" if r.get("type") == "Main material" else "Supporting", r.get("source_items")])
            for c in sh[sh.max_row]:
                c.alignment = wrap
            sh.cell(sh.max_row, 5).number_format = "#,##0.00"
        for col, w in zip("ABCDEFG", (6, 20, 60, 10, 22, 12, 40)):
            sh.column_dimensions[col].width = w
        sh.freeze_panes = "A2"
        sh.auto_filter.ref = f"A1:G{sh.max_row}"
        break
    for index, t in enumerate(tasks, 1):
        prog = _programme_rows(t)
        if not prog:
            continue
        sh = wb.create_sheet("Work Programme")
        heads = ["Item", "Work description", "Trade", "Stage", "Quantity", "Unit", "Gangs", "Duration (days)", "Start day", "End day", "Starts after item", "Critical path"]
        has_dates = "start_date" in prog[0]
        if has_dates:
            heads += ["Start date", "End date"]
        wd_ = int((t.get("engineering_result") or {}).get("programme_working_days_per_week") or 6)
        total_days = int((t.get("engineering_result") or {}).get("total_duration_days") or 0)
        plen = 1 if total_days <= 36 else wd_
        n_weeks = max(1, math.ceil(total_days / plen))
        base = len(heads)
        sh.append(heads + [(f"D{w}" if plen == 1 else f"W{w}") for w in range(1, n_weeks + 1)])
        for c in sh[1]:
            c.fill, c.font, c.alignment = header_fill, header_font, wrap
        for r in prog:
            row = [r.get("item_no"), r.get("description"), r.get("trade"), r.get("stage"), r.get("quantity"), r.get("unit"), r.get("gangs"),
                   r.get("duration_days"), r.get("start_day"), r.get("end_day"), r.get("starts_after_item"), r.get("critical_path")]
            if has_dates:
                row += [r.get("start_date"), r.get("end_date")]
            sh.append(row + [""] * n_weeks)
            w0, w1 = (int(r["start_day"]) - 1) // plen, (int(r["end_day"]) - 1) // plen
            fill = PatternFill("solid", fgColor="C0392B" if r.get("critical_path") == "Yes" else "4A78B5")
            for w in range(w0, w1 + 1):
                sh.cell(sh.max_row, base + 1 + w).fill = fill
            for c in sh[sh.max_row][:base]:
                c.alignment = wrap
        sh.append([])
        sh.append(["Estimated site duration", f"{total_days} working days (about {(t.get('engineering_result') or {}).get('total_duration_weeks')} weeks, {wd_}-day week)"])
        sh.append(["Basis", "Indicative programme for planning; red = critical path. To be confirmed with the site team."])
        for i, w in enumerate((7, 50, 16, 24, 10, 8, 7, 11, 9, 9, 12, 10) + ((12, 12) if has_dates else ())):
            sh.column_dimensions[sh.cell(1, i + 1).column_letter].width = w
        for w in range(n_weeks):
            sh.column_dimensions[sh.cell(1, base + 1 + w).column_letter].width = 4.5
        sh.freeze_panes = "C2"
        break
    styled(wb.create_sheet("Calculation Steps"), ["Task", "Step", "Calculation step", "Detail"], step_rows, [30, 8, 50, 40])
    styled(wb.create_sheet("Assumptions & Notes"), ["Task", "Type", "Detail"], note_rows, [30, 16, 100])
    cands = [(c.get("title"), c.get("authority"), c.get("edition") or "-", "Candidate - not verified") for c in governance.get("candidate_sources") or []]
    if check_rows:
        styled(wb.create_sheet("Compliance"), ["Task", "Authority", "Code / edition", "Clause", "Requirement", "Project value", "Required value", "Status"],
               check_rows, [30, 16, 24, 14, 40, 16, 16, 18])
    else:
        styled(wb.create_sheet("Compliance"), ["Standard", "Authority", "Edition", "Status"], cands or [("No compliance claim made", "", "", "")], [50, 20, 14, 28])

    sources = manifest.get("source_evidence", []) or []
    styled(wb.create_sheet("Source Evidence"), ["Type", "Attachment ID", "Filename", "MIME", "SHA-256", "Extraction Status", "Warnings"],
           [(s.get("type"), s.get("attachment_id"), s.get("filename"), s.get("mime_type"), s.get("sha256"), s.get("extraction_status"), _text(s.get("extraction_warnings"))) for s in sources],
           [14, 30, 36, 24, 66, 18, 40])

    # Technical records kept for audit (raw form).
    def add_json_sheet(title: str, payload: Any):
        sh = wb.create_sheet(title)
        if isinstance(payload, dict):
            sh.append(["Field", "Value"])
            for k, v in payload.items():
                sh.append([str(k), _text(v)])
        elif isinstance(payload, list):
            sh.append(["Index", "Value"])
            for i, v in enumerate(payload, 1):
                sh.append([i, _text(v)])
        else:
            sh.append(["Value", _text(payload)])
        sh.column_dimensions["A"].width = 28
        sh.column_dimensions["B"].width = 100
        return sh

    add_json_sheet("Manifest", manifest)
    add_json_sheet("Request", request)
    add_json_sheet("Workflow", evidence_bundle.get("workflow", {}))
    add_json_sheet("Result", evidence_bundle.get("result", {}))
    add_json_sheet("QA", evidence_bundle.get("qa", {}))
    add_json_sheet("Governance", governance)
    add_json_sheet("Review", evidence_bundle.get("review", {}))

    sh = wb.create_sheet("Tasks")
    sh.append(["Task ID", "Capability", "Status", "Objective", "Inputs", "Engineering Result", "Calculation Trace", "Assumptions", "Warnings", "Compliance"])
    for t in tasks:
        sh.append([t.get("task_id"), t.get("capability_id"), t.get("status"), t.get("objective"), _text(t.get("inputs")), _text(t.get("engineering_result")),
                   _text(t.get("calculation_trace")), _text(t.get("assumptions")), _text(t.get("warnings")), _text(t.get("compliance"))])
    for col in range(1, 11):
        sh.column_dimensions[chr(64 + col)].width = 24

    from branding import brand, footer_line
    b = brand()
    summary = wb["Summary"] if "Summary" in wb.sheetnames else wb.worksheets[0]
    summary.append([])
    summary.append([b["name"], f"{b['tagline']}  |  {b['address']}  |  {b['email']}  |  {b['phone']}"])
    summary.cell(row=summary.max_row, column=1).font = Font(bold=True, color="14233C")
    for sheet in wb.worksheets:
        sheet.oddHeader.left.text = b["name"].upper()
        sheet.oddHeader.left.size = 9
        sheet.oddHeader.right.text = "Controlled engineering record"
        sheet.oddHeader.right.size = 8
        sheet.oddFooter.left.text = footer_line()
        sheet.oddFooter.left.size = 7
        sheet.oddFooter.right.text = "Page &P of &N"
        sheet.oddFooter.right.size = 7
        sheet.sheet_properties.tabColor = "C8442B" if sheet.title == "Summary" else "14233C"
    wb.properties.creator = b["name"]
    wb.properties.title = f"Evidence workbook - {job.job_id}"

    out = io.BytesIO()
    wb.save(out)
    data = out.getvalue()
    try:  # faint Mars watermark behind every sheet; never block a report on cosmetics
        from branding.watermark import watermark_png_bytes
        from reports.xlsx_watermark import apply_background
        data = apply_background(data, watermark_png_bytes())
    except Exception:
        pass
    return data
