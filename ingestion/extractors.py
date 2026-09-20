from __future__ import annotations

import csv
import io
import json
import mimetypes
import re
import zipfile
from pathlib import Path
from typing import Any, Dict, Iterable, List, Tuple
from xml.etree import ElementTree as ET

from .models import DocumentChunk, ExtractionResult

MAX_FILE_BYTES = 25 * 1024 * 1024
MAX_ZIP_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_ZIP_MEMBERS = 3000

TEXT_EXTENSIONS = {".txt", ".md", ".rst", ".log", ".yaml", ".yml", ".json", ".xml", ".html", ".htm"}
SPREADSHEET_EXTENSIONS = {".xlsx", ".csv"}
CAD_TEXT_EXTENSIONS = {".dxf"}
CAD_BINARY_EXTENSIONS = {".dwg", ".dws", ".dwt"}


def detect_source_type(filename: str, mime_type: str | None = None) -> str:
    ext = Path(filename).suffix.lower()
    mime = (mime_type or mimetypes.guess_type(filename)[0] or "").lower()
    if ext == ".pdf" or mime == "application/pdf":
        return "pdf"
    if ext == ".docx" or "wordprocessingml" in mime:
        return "docx"
    if ext == ".xlsx" or "spreadsheetml" in mime:
        return "xlsx"
    if ext == ".csv" or mime == "text/csv":
        return "csv"
    if ext in CAD_TEXT_EXTENSIONS:
        return "dxf"
    if ext in CAD_BINARY_EXTENSIONS:
        return "cad_binary"
    if ext in TEXT_EXTENSIONS or mime.startswith("text/") or mime in {"application/json", "application/xml"}:
        return "text"
    return "binary"


def _validate_size(data: bytes) -> None:
    if len(data) > MAX_FILE_BYTES:
        raise ValueError(f"File exceeds maximum ingestion size of {MAX_FILE_BYTES // (1024 * 1024)} MB")


def _safe_zip(data: bytes) -> zipfile.ZipFile:
    zf = zipfile.ZipFile(io.BytesIO(data))
    infos = zf.infolist()
    if len(infos) > MAX_ZIP_MEMBERS:
        zf.close()
        raise ValueError("Archive contains too many members")
    uncompressed = sum(i.file_size for i in infos)
    if uncompressed > MAX_ZIP_UNCOMPRESSED_BYTES:
        zf.close()
        raise ValueError("Archive expands beyond the permitted ingestion size")
    return zf


def _chunk_text(text: str, attachment_id: str, *, max_chars: int = 4500) -> List[DocumentChunk]:
    cleaned = text.replace("\x00", "").strip()
    if not cleaned:
        return []
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", cleaned) if p.strip()]
    chunks: List[str] = []
    current = ""
    for paragraph in paragraphs:
        if len(current) + len(paragraph) + 2 <= max_chars:
            current = f"{current}\n\n{paragraph}".strip()
        else:
            if current:
                chunks.append(current)
            if len(paragraph) <= max_chars:
                current = paragraph
            else:
                for start in range(0, len(paragraph), max_chars):
                    chunks.append(paragraph[start : start + max_chars])
                current = ""
    if current:
        chunks.append(current)
    return [DocumentChunk.create(attachment_id=attachment_id, index=i, text=chunk) for i, chunk in enumerate(chunks)]


def _extract_text(data: bytes, filename: str, mime_type: str, attachment_id: str) -> ExtractionResult:
    text = data.decode("utf-8-sig", errors="replace")
    source_type = "text"
    warnings: List[str] = []
    ext = Path(filename).suffix.lower()
    if ext in {".json", ".yaml", ".yml", ".xml"}:
        try:
            if ext == ".json":
                obj = json.loads(text)
                text = json.dumps(obj, indent=2, ensure_ascii=False)
        except Exception as exc:
            warnings.append(f"Structured-text parsing failed; raw text retained: {exc}")
    chunks = _chunk_text(text, attachment_id)
    return ExtractionResult("extracted", source_type, mime_type, text, {"characters": len(text)}, warnings, [], chunks)


def _extract_csv(data: bytes, filename: str, mime_type: str, attachment_id: str) -> ExtractionResult:
    text = data.decode("utf-8-sig", errors="replace")
    reader = csv.reader(io.StringIO(text))
    rows = list(reader)
    rendered = "\n".join(" | ".join(cell.strip() for cell in row) for row in rows)
    chunks = _chunk_text(rendered, attachment_id)
    return ExtractionResult(
        "extracted",
        "csv",
        mime_type,
        rendered,
        {"rows": len(rows), "columns": max((len(r) for r in rows), default=0)},
        [],
        [],
        chunks,
    )


def _extract_xlsx(data: bytes, filename: str, mime_type: str, attachment_id: str) -> ExtractionResult:
    warnings: List[str] = []
    with _safe_zip(data) as zf:
        names = set(zf.namelist())
        shared: List[str] = []
        if "xl/sharedStrings.xml" in names:
            root = ET.fromstring(zf.read("xl/sharedStrings.xml"))
            ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
            for si in root.findall("m:si", ns):
                text = "".join(t.text or "" for t in si.findall(".//m:t", ns))
                shared.append(text)
        wb_root = ET.fromstring(zf.read("xl/workbook.xml"))
        rel_root = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
        ns_main = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
        ns_rel = "http://schemas.openxmlformats.org/package/2006/relationships"
        rels = {r.attrib["Id"]: r.attrib["Target"] for r in rel_root.findall(f"{{{ns_rel}}}Relationship")}
        sheet_records: List[Tuple[str, str]] = []
        for sheet in wb_root.findall(f"{{{ns_main}}}sheets/{{{ns_main}}}sheet"):
            name = sheet.attrib.get("name", "Sheet")
            rid = sheet.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
            target = rels.get(rid, "")
            if target.startswith("/"):
                target = target.lstrip("/")
            if not target.startswith("xl/"):
                target = "xl/" + target
            sheet_records.append((name, target))

        all_lines: List[str] = []
        sheet_meta: List[Dict[str, Any]] = []
        for sheet_name, target in sheet_records:
            if target not in names:
                warnings.append(f"Worksheet XML not found for {sheet_name}")
                continue
            root = ET.fromstring(zf.read(target))
            rows = []
            for row in root.findall(f".//{{{ns_main}}}row"):
                cells = []
                for c in row.findall(f"{{{ns_main}}}c"):
                    value = c.find(f"{{{ns_main}}}v")
                    v = value.text if value is not None else ""
                    cell_type = c.attrib.get("t")
                    if cell_type == "s" and v.isdigit() and int(v) < len(shared):
                        v = shared[int(v)]
                    elif cell_type == "inlineStr":
                        v = "".join(t.text or "" for t in c.findall(f".//{{{ns_main}}}t"))
                    cells.append(v)
                if cells:
                    rows.append(cells)
                    all_lines.append(f"[{sheet_name}] " + " | ".join(cells))
            sheet_meta.append({"name": sheet_name, "rows": len(rows)})

        text = "\n".join(all_lines)
        chunks = _chunk_text(text, attachment_id)
        return ExtractionResult(
            "extracted",
            "xlsx",
            mime_type,
            text,
            {"sheets": sheet_meta, "worksheet_count": len(sheet_meta)},
            warnings,
            [],
            chunks,
        )


def _extract_pdf(data: bytes, filename: str, mime_type: str, attachment_id: str) -> ExtractionResult:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    page_texts: List[str] = []
    warnings: List[str] = []
    for index, page in enumerate(reader.pages):
        try:
            page_texts.append(page.extract_text() or "")
        except Exception as exc:
            page_texts.append("")
            warnings.append(f"Page {index + 1}: text extraction failed: {exc}")
    text = "\n\n".join(f"[Page {i + 1}]\n{page}" for i, page in enumerate(page_texts) if page.strip())
    if not text.strip():
        warnings.append("No extractable text found; PDF may be scanned/image-only. OCR is not performed automatically.")
    chunks = _chunk_text(text, attachment_id)
    return ExtractionResult(
        "extracted" if text.strip() else "extraction_incomplete",
        "pdf",
        mime_type,
        text,
        {"pages": len(reader.pages), "pages_with_text": sum(1 for p in page_texts if p.strip())},
        warnings,
        [],
        chunks,
    )


def _extract_docx(data: bytes, filename: str, mime_type: str, attachment_id: str) -> ExtractionResult:
    from docx import Document
    doc = Document(io.BytesIO(data))
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    tables: List[str] = []
    for ti, table in enumerate(doc.tables, 1):
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells]
            tables.append(f"[Table {ti}] " + " | ".join(values))
    text = "\n\n".join(paragraphs + tables)
    chunks = _chunk_text(text, attachment_id)
    return ExtractionResult("extracted", "docx", mime_type, text, {"paragraphs": len(paragraphs), "tables": len(doc.tables)}, [], [], chunks)


def _extract_dxf_metadata(data: bytes, filename: str, mime_type: str, attachment_id: str) -> ExtractionResult:
    text = data.decode("utf-8", errors="replace")
    lines = [line.strip() for line in text.splitlines()]
    units = None
    layers: List[str] = []
    entity_counts: Dict[str, int] = {}
    in_layer_table = False
    i = 0
    while i + 1 < len(lines):
        code, value = lines[i], lines[i + 1]
        if code == "9" and value == "$INSUNITS" and i + 3 < len(lines):
            try:
                units = int(lines[i + 3])
            except Exception:
                pass
        if value == "LAYER" and code == "0":
            in_layer_table = True
        elif in_layer_table and code == "2" and value not in {"ENDTAB", "LAYER"}:
            if value not in layers:
                layers.append(value)
        if code == "0" and value not in {"SECTION", "ENDSEC", "TABLE", "ENDTAB", "LAYER", "EOF"}:
            entity_counts[value] = entity_counts.get(value, 0) + 1
        if code == "0" and value == "ENDTAB":
            in_layer_table = False
        i += 2
    summary = f"DXF metadata: units_code={units}, layers={len(layers)}, entities={sum(entity_counts.values())}."
    metadata = {"units_code": units, "layers": layers, "entity_counts": entity_counts, "binary_geometry_extraction": False}
    warnings = ["DXF geometry is not converted into engineering geometry in this V1 extractor; only file metadata is extracted."]
    return ExtractionResult("metadata_only", "dxf", mime_type, summary, metadata, warnings, [], [DocumentChunk.create(attachment_id=attachment_id, index=0, text=summary, metadata=metadata)])


def _extract_binary_cad(data: bytes, filename: str, mime_type: str, attachment_id: str) -> ExtractionResult:
    return ExtractionResult(
        "metadata_only",
        "cad_binary",
        mime_type,
        "",
        {"format": Path(filename).suffix.lower().lstrip("."), "geometry_extraction": False, "bytes": len(data)},
        ["Binary CAD geometry extraction is not performed by this V1 layer. Register the file for downstream CAD tooling or human review."],
        [],
        [],
    )


def extract_bytes(data: bytes, *, filename: str, mime_type: str | None = None, attachment_id: str = "attachment") -> ExtractionResult:
    _validate_size(data)
    mime = mime_type or mimetypes.guess_type(filename)[0] or "application/octet-stream"
    source_type = detect_source_type(filename, mime)
    try:
        if source_type == "pdf":
            return _extract_pdf(data, filename, mime, attachment_id)
        if source_type == "docx":
            return _extract_docx(data, filename, mime, attachment_id)
        if source_type == "xlsx":
            return _extract_xlsx(data, filename, mime, attachment_id)
        if source_type == "csv":
            return _extract_csv(data, filename, mime, attachment_id)
        if source_type == "text":
            return _extract_text(data, filename, mime, attachment_id)
        if source_type == "dxf":
            return _extract_dxf_metadata(data, filename, mime, attachment_id)
        if source_type == "cad_binary":
            return _extract_binary_cad(data, filename, mime, attachment_id)
        return ExtractionResult("unsupported", source_type, mime, warnings=[f"No extractor registered for {filename}"])
    except Exception as exc:
        return ExtractionResult("failed", source_type, mime, errors=[str(exc)])
